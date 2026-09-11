#!/usr/bin/env python
"""Builds module_summary.docx in the Udacity APA 7 template (Rubrics/APA7_template.docx) with the sections
the Module 02 instructions prescribe, then converts it to module_summary.pdf with Word (COM automation
through PowerShell). Run:  python make_summary.py
Every number in the text comes from the notebook outputs / SUMMARY files; citations are APA author–year
and the References list contains exactly the sources cited."""
import subprocess, shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parents[3] / "Rubrics" / "APA7_template.docx"
FIG = HERE / "notebook" / "figures"
OUT_DOCX = HERE / "module_summary.docx"

doc = Document(str(TEMPLATE))
body = doc.element.body
for p in list(doc.paragraphs):           # keep the template's styles, section and page setup; drop its placeholder text
    p._element.getparent().remove(p._element)

def para(text="", bold=False, center=False, italic=False, size=None, indent_first=True):
    p = doc.add_paragraph()
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not indent_first: p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def heading(text, level=1):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    if level == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = True
    if level == 2: r.italic = False
    return p

def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def figure(name, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_picture(str(FIG / name), width=Inches(6.0))
    c = doc.add_paragraph(); c.paragraph_format.first_line_indent = Pt(0)
    r = c.add_run(caption); r.italic = True; r.font.size = Pt(10)

TITLE = "The Opponent Atlas: A Reproducible Data Workflow on NASA's PAHdb Computed Library"

# ---------------- title page
for _ in range(6): para(indent_first=False)
para(TITLE, bold=True, center=True, indent_first=False)
para(indent_first=False)
para("Frederic Petrignani", center=True, indent_first=False)
para("Udacity AI Mastery Capstone", center=True, indent_first=False)
para("Module 02 — AI Programming Foundations Project", center=True, indent_first=False)
para("Written summary (module_summary.pdf)", center=True, indent_first=False)
para("11 September 2026", center=True, indent_first=False)
page_break()

# ---------------- body
para(TITLE, bold=True, center=True, indent_first=False)

heading("Overview")
para("I built a reproducible data workflow that reads the public NASA Ames PAH IR Spectroscopic Database "
     "(PAHdb), computed library version 4.00, into tidy tables and explores what it contains. The dataset is "
     "the library's XML download (https://www.astrochemistry.org/pahdb/theoretical/4.00), parsed into a "
     "species table of 10,749 rows and 27 columns and a band table of 2,517,399 rows. The library is the "
     "\"opponent\" of my larger capstone project: the reference predictions that a later, more accurate "
     "pipeline will be compared against. The data are computed science data, not AI-generated (Ricca et al., 2026).")

heading("Dataset Description")
para("PAHdb is a library of computed infrared spectra of polycyclic aromatic hydrocarbons and related "
     "species, maintained at NASA Ames (Boersma et al., 2014; Ricca et al., 2026). Each species record carries "
     "a formula, an electric charge, element counts, the point-group symmetry, the quantum-chemical method and "
     "basis set (written only in a free-text route line), and a list of vibrational transitions, each with a "
     "frequency, an intensity in km/mol and the scale factor applied to it. The species table I derived has "
     "10,749 rows (4,479 neutral molecules, 2,162 cations, 2,868 dications, 1,231 anions and 9 trications) and "
     "27 columns; the band table has 2,517,399 rows. The key variables I focused on were the carbon count, the "
     "charge, the basis set, the stored scale factor, the band position and intensity, and a frequency-range "
     "family label that I added for exploration only.")

heading("Workflow Description")
para("Ingestion: a streaming XML parser (Python's iterparse) reads the 503 MB unpacked file in a few minutes "
     "and writes one row per species and one row per band, together with the SHA-256 checksum of the source "
     "file and the XML root attributes, so that the version is pinned by content rather than by file name. "
     "The notebook then loads the species table with pandas and displays its first rows. Cleaning: two "
     "documented functions separate the charge suffix from the formula string and flag species whose basis set "
     "cannot be resolved from the route line. Exploratory analysis: one documented function counts species per "
     "rung of my project's test ladder for each of five comparison lines. Visualizations: five titled and "
     "labelled figures. Summary: findings, assumptions and limitations in the notebook's final section. The "
     "whole notebook is rebuilt and executed by a script, so a reviewer obtains the same figures from the same "
     "tables; this follows the guidance that a data-science workflow should be re-runnable from its inputs "
     "without manual steps (Danchev, 2022).")

heading("Key Decisions and Assumptions")
para("Cleaning choices. The formula field encodes charge as a suffix (for example C54H18- or C384H48+2), so "
     "6,270 of the 10,749 formula strings would have split one composition across several labels; I added a "
     "formula_core column and verified that the suffix agrees with the charge column for every row. The basis "
     "set is parsed from a Gaussian route comment; 10,703 species resolve to 6-31G*, 14 to 4-31G, and 32 carry "
     "chkbas (basis read from a checkpoint), which I flag rather than guess, because assigning them to the "
     "majority basis would move the boundary at which the coarser basis begins. Species with the same formula "
     "and charge but different identifiers are isomers, not duplicates (1,159 groups, 638 with more than one "
     "identifier, the largest with 358), so they are kept; dropping them would remove 9,590 real molecules. "
     "These are exactly the kinds of silent transformations that reproducible-workflow practice asks to be made "
     "explicit and testable (Danchev, 2022).")
para("What I focused on in the exploration. Coverage by size and charge, the boundary between basis sets, the "
     "scale factors as actually stored, and coverage of my ladder of test molecules by each comparison line. "
     "Figure 1 was designed to show the size and charge distribution and the basis boundary; Figure 2 the "
     "stored scale factors against frequency, with the factors the current paper describes drawn as reference "
     "lines; Figure 3 the intensity-weighted band positions and their family labels; Figure 4 which test "
     "molecules have entries in which library; Figure 5 the largest molecule's computed spectrum.")

heading("Results and Interpretation")
figure("fig1_species_by_size_charge.png", "Figure 1. PAHdb theoretical v4.00: 10,749 species by carbon count and charge; the dashed line marks where the 4-31G basis begins (212 carbons).")
para("Most of the library lies between 20 and 100 carbon atoms (Figure 1). The coarser 4-31G basis appears "
     "only from 212 carbons upward, for 14 species; the largest 6-31G* species has 294 carbons. The 101–386-carbon "
     "bin holds 774 species, the number the current paper gives (Ricca et al., 2026).")
figure("fig2_scale_factors_as_stored.png", "Figure 2. Scale factors stored with the bands, against unscaled frequency (six most frequent factors; 4,000-band samples); the dotted lines are the factors the version 4.00 paper describes.")
para("Figure 2 shows the first surprise. The current paper states three scale factors, 0.964 for the C–H "
     "stretches near 3 µm, 0.979 for 4–9 µm and 0.975 beyond 9 µm, fitted to 25 gas-phase laboratory bands "
     "(Ricca et al., 2026). None of the 2,517,399 stored bands carries any of these values. The 6-31G* species "
     "carry 0.9794, 0.9691 and 0.9597, and the 4-31G species 0.9563, 0.9523 and 0.9595, in three frequency "
     "regions; these are the version 3.00 factors published in the earlier paper's Table 2 (Bauschlicher et "
     "al., 2018). The difference is 4–15 cm⁻¹ at the band positions. I read the stored frequencies as scaled "
     "values because dividing them by the stored factor returns numbers in the range of raw DFT results, "
     "checked on naphthalene's C–H stretch. My project recorded, before any result of its own existed, that it "
     "scores against the library as served and prints a second column at the paper's factors.")
figure("fig3_positions_by_family.png", "Figure 3. Neutral species: intensity-weighted band positions (left) and summed intensity per frequency-range family label (right).")
para("Figure 3 shows the familiar PAH pattern: the 7.7 µm C–C/C–H in-plane region carries the most summed "
     "intensity, followed by the C–H stretch, the C–H out-of-plane bands, the 8.6 µm bend and the 6.2 µm C–C "
     "stretch. The family label is a frequency-range rule, not a mode assignment.")
figure("fig4_ladder_coverage.png", "Figure 4. Species per ladder rung for each comparison line (all charges); a dash means no entry.")
para("Figure 4 shows the second surprise: benzene, the project's first test molecule, is in no library except "
     "the small anharmonic one. The smallest entries of the computed library are phenol- and indene-like "
     "species; the smallest plain PAH is naphthalene. From naphthalene upward all five lines have entries; "
     "above coronene only the computed library and the machine-learning simulation of Mai et al. (2025) "
     "remain, and above 216 carbons only the computed library.")
figure("fig5_c384h48_sticks.png", "Figure 5. C384H48 (uid 617): the scaled harmonic stick spectrum at B3LYP/4-31G, 1,290 modes.")
para("Figure 5 answers an open question of my project: the target class of the largest molecules exists in "
     "the library. C384H48 is present as a neutral molecule (uid 617) and as a dication (uid 4447), and the "
     "scaled-DFT spectrum is the only prediction of any kind at that size.")

heading("Responsible Practice (Bias and Data Quality)")
para("Three cleaning decisions could have biased the picture. Leaving the charge suffix in the formula would "
     "have undercounted neutral species and overcounted distinct compositions; assigning the 32 unresolved "
     "species to the majority basis would have misplaced the basis boundary; and treating isomers as "
     "duplicates would have removed 9,590 molecules and misstated the coverage of the larger sizes. I made each "
     "step explicit, asserted the charge check on every row, and kept unresolved cases flagged. A separate "
     "quality question is the disagreement between the file and its describing paper on the scale factors; I "
     "report it as a version-mismatch finding, have not asked the database team for its cause, and avoid "
     "choosing the version that would flatter my project's later results. Handling such choices in the open, "
     "with the code and the checksums beside the claims, is what makes a workflow trustworthy rather than merely "
     "repeatable (Danchev, 2022).")

heading("Reproducibility")
para("The repository contains the parser, the two smaller readers for the other comparison lines, the notebook "
     "builder, the derived tables (the species table and small band tables are committed; the 51 MB band "
     "table is regenerated by one command) and a requirements.txt produced with pip freeze. Someone else "
     "installs the requirements, places the PAHdb download in data/, runs the parser, and executes "
     "notebook/data_workflow.ipynb top to bottom, or runs notebook/make_notebook.py to rebuild and execute it. "
     "Every output file records the checksum of its input. The work was developed on a branch "
     "(module-02-opponent-atlas) with many small commits and merged into master; the raw downloads stay out of "
     "version control and are identified by hash. This mirrors the practices of environment specification, "
     "version control and scripted execution that Danchev (2022) sets out for reproducible data science.")

heading("Sources and Citations")
para("Sources are cited where I define a concept, describe the database or justify a practice; my own counts, "
     "figures and observations are not cited. The two required source types are present: the assigned "
     "peer-reviewed article on reproducible data science (Danchev, 2022) and peer-reviewed articles describing "
     "the dataset and its scale factors (Bauschlicher et al., 2018; Ricca et al., 2026), with the database's "
     "requested citations (Boersma et al., 2014; Mattioda et al., 2020) and the two other comparison lines "
     "(Bos et al., 2025; Mai et al., 2025).")
para("Use of AI assistance. This module was planned, coded and written with Claude (Anthropic) as an "
     "assistant under my direction; every download, run and decision is mine, every number traces to a "
     "named script and commit in the repository, and the text was revised against the actual outputs.")

page_break()
heading("References")
REFS = [
 "Bauschlicher, C. W., Jr., Ricca, A., Boersma, C., & Allamandola, L. J. (2018). The NASA Ames PAH IR Spectroscopic Database: Computational version 3.00 with updated content and the introduction of multiple scaling factors. The Astrophysical Journal Supplement Series, 234(2), 32. https://doi.org/10.3847/1538-4365/aaa019",
 "Boersma, C., Bauschlicher, C. W., Jr., Ricca, A., Mattioda, A. L., Cami, J., Peeters, E., Sánchez de Armas, F., Puerta Saborido, G., Hudgins, D. M., & Allamandola, L. J. (2014). The NASA Ames PAH IR Spectroscopic Database version 2.00: Updated content, web site, and on(off)line tools. The Astrophysical Journal Supplement Series, 211(1), 8. https://doi.org/10.1088/0067-0049/211/1/8",
 "Bos, R., King, M., Calangian, A. J., Davin, K. A., McCraley, S., Putnam, R. A., Manjarrés, J., & Stoneburner, S. J. (2025). Ethereal AI: Infrared spectra of polycyclic aromatic hydrocarbons with machine learning DFT scaling factors. ACS Omega, 10(50), 62282–62290. https://doi.org/10.1021/acsomega.5c10225",
 "Danchev, V. (2022). Reproducible data science with Python: An open learning resource. Journal of Open Source Education, 5(56), 156. https://doi.org/10.21105/jose.00156",
 "Mai, X., Wang, Z., Pan, L., Schörghuber, J., Kovács, P., Carrete, J., & Madsen, G. K. H. (2025). Computing anharmonic infrared spectra of polycyclic aromatic hydrocarbons using machine learning molecular dynamics. Monthly Notices of the Royal Astronomical Society, 541(4), 3073–3080. https://doi.org/10.1093/mnras/staf1156",
 "Mattioda, A. L., Hudgins, D. M., Boersma, C., Bauschlicher, C. W., Jr., Ricca, A., Cami, J., Peeters, E., Sánchez de Armas, F., Puerta Saborido, G., & Allamandola, L. J. (2020). The NASA Ames PAH IR Spectroscopic Database: The laboratory spectra. The Astrophysical Journal Supplement Series, 251(2), 22. https://doi.org/10.3847/1538-4365/abc2c8",
 "Ricca, A., Boersma, C., Maragkoudakis, A., Roser, J. E., Shannon, M. J., Allamandola, L. J., & Bauschlicher, C. W., Jr. (2026). The NASA Ames PAH IR Spectroscopic Database: Computational version 4.00, software tools, website, and documentation. The Astrophysical Journal Supplement Series, 282(1), 7. https://doi.org/10.3847/1538-4365/ae1c38",
]
for ref in REFS:
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Inches(-0.5); p.paragraph_format.left_indent = Inches(0.5)
    p.add_run(ref)

doc.save(str(OUT_DOCX)); print("written", OUT_DOCX)

# ---------------- PDF via Word (COM through PowerShell)
ps = f'''$w = New-Object -ComObject Word.Application; $w.Visible = $false
$d = $w.Documents.Open("{OUT_DOCX}"); $d.SaveAs2("{OUT_DOCX.with_suffix('.pdf')}", 17); $d.Close(); $w.Quit()'''
r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
print("pdf:", OUT_DOCX.with_suffix(".pdf").exists(), r.stderr[:300])
